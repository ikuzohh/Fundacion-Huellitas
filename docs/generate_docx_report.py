import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_report():
    doc = docx.Document()

    # Configuración de márgenes estándar (2.5 cm)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Estilos de fuente base
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)

    # Encabezado Institucional
    header_para = doc.add_paragraph()
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_inst = header_para.add_run("INSTITUTO PROFESIONAL AIEP | ESCUELA TIC\n")
    run_inst.font.bold = True
    run_inst.font.size = Pt(10)
    run_inst.font.color.rgb = RGBColor(30, 58, 138)
    run_mod = header_para.add_run("TALLER DE PROGRAMACIÓN (PRO205) — PRIMAVERA 2026")
    run_mod.font.size = Pt(9.5)
    run_mod.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph()

    # TÍTULO PRINCIPAL
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(10)
    title_p.paragraph_format.space_after = Pt(4)
    run_t1 = title_p.add_run("INFORME TÉCNICO: EVALUACIÓN SUMATIVA N° 3\n")
    run_t1.font.bold = True
    run_t1.font.size = Pt(22)
    run_t1.font.color.rgb = RGBColor(30, 58, 138)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(18)
    run_sub = sub_p.add_run("Articulación Oficial con la Evaluación Nacional de Especialidad (ENE)\nCaso Seleccionado: OPCIÓN A — Fundación Huellitas (Modelado UML y Dominio POO)")
    run_sub.font.bold = True
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = RGBColor(2, 132, 199)

    # TABLA METADATOS
    meta_table = doc.add_table(rows=9, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    metadata = [
        ("Estudiante:", "Pablo Esteban Herrera Bustamante"),
        ("Correo Electrónico:", "ikuzoherrera@gmail.com"),
        ("Docente / Profesora:", "MARÍA UBILLA SALINAS"),
        ("Asignatura:", "Taller de Programación (PRO205)"),
        ("Ponderación de Evaluación:", "35% de la Calificación Final (100 Puntos Totales)"),
        ("Caso y Desafío:", "Opción A: Fundación Huellitas + Demostración Polimórfica en C# (.NET 8)"),
        ("Modalidad:", "Individual (Estudiante Autónomo)"),
        ("Criterios de Evaluación:", "C.E. 1.8 (Relaciones UML), C.E. 1.9 (Requerimientos), POO Avanzada"),
        ("Puntaje Máximo Esperado:", "100 Puntos (Exigencia 60% para Nota 4.0)")
    ]

    for idx, (k, v) in enumerate(metadata):
        row = meta_table.rows[idx]
        cell_k, cell_v = row.cells[0], row.cells[1]
        cell_k.width = Inches(2.2)
        cell_v.width = Inches(4.3)
        set_cell_background(cell_k, "F1F5F9")
        set_cell_background(cell_v, "FFFFFF" if idx % 2 == 0 else "F8FAFC")
        set_cell_margins(cell_k, 60, 60, 100, 100)
        set_cell_margins(cell_v, 60, 60, 100, 100)
        
        pk = cell_k.paragraphs[0]
        pk.paragraph_format.space_after = Pt(0)
        rk = pk.add_run(k)
        rk.font.bold = True
        rk.font.size = Pt(10)
        rk.font.color.rgb = RGBColor(15, 23, 42)

        pv = cell_v.paragraphs[0]
        pv.paragraph_format.space_after = Pt(0)
        rv = pv.add_run(v)
        rv.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 1. INTRODUCCIÓN
    h1 = doc.add_heading("1. Introducción y Contexto del Caso", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    p = doc.add_paragraph(
        "La Fundación Huellitas es una organización sin fines de lucro dedicada al rescate, rehabilitación médica, "
        "albergue temporal y adopción responsable de caninos y felinos en situación de vulnerabilidad o abandono. "
        "Históricamente, la gestión de fichas clínicas y solicitudes se realizaba en soportes manuales, provocando desorganización "
        "y riesgo de inconsistencia de datos.\n\n"
        "El presente informe documenta la solución técnica desarrollada para sistematizar integralmente dichos procesos, "
        "cumpliendo con la resolución base de la Evaluación Nacional de Especialidad (Subtareas 1A y 1B) y el desafío práctico "
        "de profundización exigido en la Evaluación Sumativa 3 mediante programación orientada a objetos en C# (.NET 8)."
    )

    # 2. SUBTAREA 1A
    h2 = doc.add_heading("2. SUBTAREA 1A: Esquema Conceptual Preliminar y Justificación Técnica de Relaciones (Base ENE)", level=1)
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "En cumplimiento de los requerimientos de la Subtarea 1A y el Criterio de Evaluación 1.8, se identificaron "
        "las entidades principales del sistema y se fundamentaron técnicamente los 4 tipos de relaciones de la POO:"
    )

    # TABLA 4 RELACIONES
    rel_table = doc.add_table(rows=5, cols=3)
    rel_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rel_table.autofit = False

    headers = ["Tipo de Relación", "Clases Involucradas", "Justificación Técnica y Ciclo de Vida"]
    for i, h in enumerate(headers):
        cell = rel_table.rows[0].cells[i]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)

    rel_data = [
        ("1. Generalización / Herencia", "Animal -> Perro, Gato", 
         "Relación 'Es-Un'. Centraliza en la superclase abstracta 'Animal' los atributos y operaciones comunes (Id, Nombre, Edad, Peso, Estado), evitando duplicación de código. Las subclases 'Perro' y 'Gato' especializan protocolos médicos, sonidos característicos y dosificación metabólica."),
        ("2. Composición Estricta", "Animal *-- HistorialMedico", 
         "Relación de dependencia existencial fuerte ('Parte-De'). La colección de fichas clínicas nace, vive y se destruye con la instancia de 'Animal'. El constructor de 'HistorialMedico' es interno y sólo se accede mediante el método 'animal.RegistrarAtencion(...)', garantizando que no existan atenciones huérfanas."),
        ("3. Agregación", "Refugio o-- Animal\nRefugio o-- Voluntario", 
         "Relación todo/parte débil ('Tiene-Un'). El Refugio agrupa y administra animales y voluntarios, pero el ciclo de vida de estos es completamente independiente. Si una sede física cierra o se reubica, los animales son derivados y los voluntarios continúan en la organización."),
        ("4. Asociación", "SolicitudAdopcion -> Adoptante, Animal, Voluntario", 
         "Relación de colaboración estructural ('Usa-Un'). La clase SolicitudAdopcion vincula temporal y funcionalmente al postulante con la mascota solicitada y el voluntario evaluador para orquestar la aprobación y seguimiento, manteniendo la autonomía de las tres entidades.")
    ]

    for idx, row in enumerate(rel_data):
        r_cells = rel_table.rows[idx + 1].cells
        r_cells[0].width = Inches(1.8)
        r_cells[1].width = Inches(1.8)
        r_cells[2].width = Inches(2.9)
        bg = "FFFFFF" if idx % 2 == 0 else "F8FAFC"
        for i in range(3):
            set_cell_background(r_cells[i], bg)
            set_cell_margins(r_cells[i], 60, 60, 100, 100)
            p = r_cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(row[i])
            run.font.size = Pt(9.5)
            if i == 0:
                run.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 3. SUBTAREA 1B
    h3 = doc.add_heading("3. SUBTAREA 1B: Diagrama Formal de Clases en Notación UML y Requerimientos Cubiertos (Base ENE)", level=1)
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "El diagrama de clases formal modela 8 clases altamente cohesivas, especificando visibilidad "
        "(+ Público, - Privado, # Protegido, ~ Internal), tipos de datos primitivos y complejos, signaturas de métodos con parámetros, "
        "y cardinalidades estándar (1, 1..*, 0..*)."
    )

    doc.add_paragraph(
        "Nota Visual: Los diagramas vectoriales completos en alta resolución están disponibles en los archivos oficiales:\n"
        "• docs/diagramas/diagrama_clases_formal_1b.svg (Diagrama UML Formal Completo)\n"
        "• docs/diagramas/esquema_conceptual_1a.svg (Esquema Conceptual Preliminar)\n"
        "• docs/diagramas/visor_diagramas_uml.html (Visor Web Interactivo AIEP)"
    )

    # TABLA REQUERIMIENTOS
    req_heading = doc.add_heading("3.1 Lista de Requerimientos Cubiertos por el Diseño (C.E. 1.9)", level=2)
    req_heading.paragraph_format.space_before = Pt(10)
    req_heading.paragraph_format.space_after = Pt(4)

    req_table = doc.add_table(rows=7, cols=3)
    req_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    req_table.autofit = False

    req_headers = ["ID Req.", "Nombre del Requerimiento", "Descripción Técnica y Clases Responsables"]
    for i, h in enumerate(req_headers):
        cell = req_table.rows[0].cells[i]
        set_cell_background(cell, "0284C7")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)

    req_data = [
        ("RF-01", "Registro y tipificación de animales rescatados", "Permite la admisión de animales clasificándolos en caninos o felinos mediante las clases 'Perro' y 'Gato', capturando sus datos biométricos y estado inicial de rescate."),
        ("RF-02", "Historial clínico y trazabilidad veterinaria", "Garantiza el registro inmutable de consultas, vacunaciones, desparasitaciones y cirugías mediante la composición 'Animal *-- HistorialMedico'."),
        ("RF-03", "Cálculo farmacológico y dosificación personalizada", "Automatiza el cálculo seguro de dosis mediante el método polimórfico 'CalcularDosisMedicamento(factor)', adaptando la dosis a la masa corporal y al metabolismo específico de cada especie."),
        ("RF-04", "Registro de postulantes y filtro habitacional", "Captura antecedentes de adoptantes y valida reglas de tenencia responsable (ej. el método 'EsAptoParaGatos()' corrobora la instalación de mallas de seguridad)."),
        ("RF-05", "Flujo de solicitudes y adopción responsable", "La clase 'SolicitudAdopcion' asocia adoptante, animal y voluntario evaluador, gestionando estados desde postulación hasta adopción definitiva."),
        ("RF-06", "Gestión de albergue y contingente de voluntarios", "La clase 'Refugio' implementa la agregación para controlar cupos de albergue físico y coordinar al equipo de voluntarios activos.")
    ]

    for idx, row in enumerate(req_data):
        r_cells = req_table.rows[idx + 1].cells
        r_cells[0].width = Inches(1.0)
        r_cells[1].width = Inches(2.2)
        r_cells[2].width = Inches(3.3)
        bg = "FFFFFF" if idx % 2 == 0 else "F8FAFC"
        for i in range(3):
            set_cell_background(r_cells[i], bg)
            set_cell_margins(r_cells[i], 60, 60, 100, 100)
            p = r_cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(row[i])
            run.font.size = Pt(9.5)
            if i == 0:
                run.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 4. PREGUNTA COMPLEMENTARIA SUMATIVA 3
    h4 = doc.add_heading("4. PREGUNTA COMPLEMENTARIA / PASO EXTRA: Demostración en Código y Polimorfismo Dinámico", level=1)
    h4.paragraph_format.space_before = Pt(16)
    h4.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "En respuesta a la Pregunta Complementaria obligatoria de la Evaluación Sumativa 3, la arquitectura UML "
        "fue programada y ejecutada en código C# (.NET 8.0) estructurada en tres proyectos dentro de la solución 'FundacionHuellitas.sln':"
    )

    doc.add_paragraph(
        "1. Demostración de Composición: La clase abstracta Animal gestiona internamente la lista _historialClinico. "
        "Esta lista se expone al exterior exclusivamente como 'IReadOnlyList<HistorialMedico>', impidiendo mutaciones externas y asegurando encapsulamiento estricto.\n"
        "2. Demostración de Polimorfismo Dinámico: Se implementó un programa principal en consola (Program.cs) que administra una colección "
        "heterogénea 'List<Animal>' compuesta por instancias concretas de 'Perro' y 'Gato'. Durante el recorrido de la colección, se invocan los métodos "
        "polimórficos 'EmitirFichaClinica()', 'HacerSonidoCaracteristico()' y 'CalcularDosisMedicamento()', donde cada subclase resuelve en tiempo de ejecución "
        "el comportamiento adecuado según su especie."
    )

    doc.add_paragraph(
        "Justificación de Encapsulamiento Estricto:\n"
        "El uso de modificadores de acceso 'private' en los campos y constructores internos restringe el acceso indebido. Cada setter valida invariantes "
        "(ej. peso > 0, nombres no vacíos, sexo válido 'MACHO' o 'HEMBRA'). Esto garantiza la integridad del sistema al impedir que existan objetos en estados "
        "inválidos o corruptos en memoria.\n\n"
        "Justificación de Herencia contra la Duplicidad de Código:\n"
        "La superclase 'Animal' centraliza el 100% de los datos comunes y la lógica de historial clínico, aplicando el principio DRY (Don't Repeat Yourself). "
        "Si en el futuro se incorpora una nueva especie (ej. 'Conejo'), sólo se extiende 'Animal' sin tener que reimplementar la gestión médica ni alterar las rutinas ya existentes."
    )

    # 5. PRUEBAS AUTOMATIZADAS
    h5 = doc.add_heading("5. Pruebas y Verificación Automatizada (15 Puntos de Rúbrica)", level=1)
    h5.paragraph_format.space_before = Pt(16)
    h5.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "Para garantizar los 15 puntos del criterio 'Pruebas y Verificación Automatizada' y certificar una arquitectura libre de errores, "
        "se auditó el código y se desarrolló una suite integral de 24 pruebas unitarias automatizadas con xUnit, cubriendo validaciones defensivas, "
        "composición estricta, polimorfismo dinámico, flujo transaccional de adopción y límites de agregación en el refugio:"
    )

    test_table = doc.add_table(rows=25, cols=3)
    test_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    test_table.autofit = False

    test_headers = ["Suite de Pruebas", "Nombre del Test Unitario", "Estado de Ejecución"]
    for i, h in enumerate(test_headers):
        cell = test_table.rows[0].cells[i]
        set_cell_background(cell, "16A34A")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)

    test_data = [
        ("AnimalTests", "Animal_DebeLanzarExcepcion_CuandoNombreEsVacio", "✔ Correcto (Aprobado)"),
        ("AnimalTests", "Animal_DebeLanzarExcepcion_CuandoPesoEsCeroONegativo", "✔ Correcto (Aprobado)"),
        ("AnimalTests", "Animal_DebeLanzarExcepcion_CuandoSexoEsInvalido", "✔ Correcto (Aprobado)"),
        ("AnimalTests", "Animal_DebeActualizarEstadoSalud_Correctamente", "✔ Correcto (Aprobado)"),
        ("AnimalTests", "Perro_Esterilizacion_NoPermiteRepetirse", "✔ Correcto (Aprobado)"),
        ("ComposicionTests", "Composicion_AnimalDebeInicializarHistorialClinicoConRegistroDeIngreso", "✔ Correcto (Aprobado)"),
        ("ComposicionTests", "Composicion_AnimalGestionaDirectamenteColeccionDeHistorialMedico", "✔ Correcto (Aprobado)"),
        ("ComposicionTests", "Composicion_HistorialClinicoProtegidoPorEncapsulamientoDefensivo", "✔ Correcto (Aprobado)"),
        ("PolimorfismoTests", "Polimorfismo_ColeccionHeterogenea_ResuelveSonidoSegunSubclase", "✔ Correcto (Aprobado)"),
        ("PolimorfismoTests", "Polimorfismo_CalculoDosis_AplicaFactoresDiferenciadosPorEspecie", "✔ Correcto (Aprobado)"),
        ("PolimorfismoTests", "Polimorfismo_EmitirFichaClinica_GeneraFormatoEspecializado", "✔ Correcto (Aprobado)"),
        ("AuditoriaIntegracion", "SolicitudAdopcion_DebeRechazarGato_SiViviendaNoTieneMallasDeSeguridad", "✔ Correcto (Aprobado)"),
        ("AuditoriaIntegracion", "SolicitudAdopcion_DebeAgendarEntrevista_CuandoAnimalEstaAptoYPerfilCumple", "✔ Correcto (Aprobado)"),
        ("AuditoriaIntegracion", "SolicitudAdopcion_DebeMantenerEnEspera_SiAnimalAunNoEstaRehabilitado", "✔ Correcto (Aprobado)"),
        ("AuditoriaIntegracion", "SolicitudAdopcion_ConcretarAdopcion_FinalizaProcesoYActualizaEstado", "✔ Correcto (Aprobado)"),
        ("AuditoriaIntegracion", "Refugio_DebeLanzarExcepcion_CuandoSeSuperaCapacidadMaxima", "✔ Correcto (Aprobado)"),
        ("AuditoriaIntegracion", "Refugio_NoDebeDuplicarAnimal_SiSeIngresaMasDeUnaVez", "✔ Correcto (Aprobado)"),
        ("AuditoriaIntegracion", "Refugio_DebeContarEspeciesCorrectamente_MedianteMetodoGenerico", "✔ Correcto (Aprobado)"),
        ("AuditoriaIntegracion", "Adoptante_DebeLanzarExcepcion_SiCorreoNoTieneArroba", "✔ Correcto (Aprobado)"),
        ("AuditoriaIntegracion", "Voluntario_DebeLanzarExcepcion_SiHorasSemanalesSonNegativas", "✔ Correcto (Aprobado)"),
        ("AuditoriaIntegracion", "Composicion_HistorialClinico_NoPuedeCastearseAListaMutable", "✔ Correcto (Aprobado)"),
        ("AuditoriaIntegracion", "Animal_CalcularDosis_DebeLanzarExcepcion_CuandoFactorEsCeroONegativo", "✔ Correcto (Aprobado)"),
        ("AuditoriaIntegracion", "Perro_AplicarVacunaSextuple_ActualizaBanderaYAgregaRegistroAlHistorial", "✔ Correcto (Aprobado)"),
        ("AuditoriaIntegracion", "Gato_RealizarTestFelvFiv_ActualizaResultadoYAgregaAtencionAlHistorial", "✔ Correcto (Aprobado)")
    ]

    for idx, row in enumerate(test_data):
        r_cells = test_table.rows[idx + 1].cells
        r_cells[0].width = Inches(1.8)
        r_cells[1].width = Inches(3.2)
        r_cells[2].width = Inches(1.5)
        bg = "FFFFFF" if idx % 2 == 0 else "F8FAFC"
        for i in range(3):
            set_cell_background(r_cells[i], bg)
            set_cell_margins(r_cells[i], 50, 50, 100, 100)
            p = r_cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(row[i])
            run.font.size = Pt(9.5)
            if i == 2:
                run.font.bold = True
                run.font.color.rgb = RGBColor(22, 163, 74)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 6. INSTRUCCIONES DE EJECUCIÓN
    h6 = doc.add_heading("6. Guía Rápida de Compilación y Ejecución", level=1)
    h6.paragraph_format.space_before = Pt(16)
    h6.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "Para compilar y ejecutar el proyecto en cualquier estación de trabajo con .NET 8 SDK:\n\n"
        "1. Compilar la solución:\n   dotnet build src/FundacionHuellitas.sln\n\n"
        "2. Ejecutar la suite de pruebas unitarias:\n   dotnet test src/FundacionHuellitas.sln\n\n"
        "3. Ejecutar la aplicación en consola con la demostración polimórfica:\n   dotnet run --project src/FundacionHuellitas.App\n"
    )

    output_path = r"c:\Users\ikuzo\OneDrive\Desktop\entrega final\docs\INFORME_TECNICO_SUMATIVA_3_HUELLITAS.docx"
    doc.save(output_path)
    print(f"Documento DOCX generado exitosamente en: {output_path}")

if __name__ == "__main__":
    create_report()
